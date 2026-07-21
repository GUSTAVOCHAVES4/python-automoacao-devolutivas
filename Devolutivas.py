import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk, simpledialog
import pandas as pd
import os
import warnings
import json
import math
import difflib
from openpyxl.styles import Alignment
from openpyxl import load_workbook
from datetime import datetime

# Bibliotecas Visuais e Leitura de PDF
HAS_PDF_VISION = True
try:
    import fitz  # PyMuPDF
    from PIL import Image, ImageTk, ImageDraw
except ImportError:
    HAS_PDF_VISION = False

# Bibliotecas do Word e PowerPoint
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pptx
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

# Biblioteca de OCR (Leitura de Escaneados)
HAS_OCR = True
try:
    import pytesseract
    # ATENÇÃO: Verifique se este é o local exato onde você instalou o Tesseract!
    pytesseract.pytesseract.tesseract_cmd = r'C:\Users\X458530\Tesseract\tesseract.exe'
except ImportError:
    HAS_OCR = False

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')


def normalizar_texto_comparacao(valor):
    """Normaliza células do Excel para comparar conteúdo sem diferença de maiúsculas/espaços."""
    if valor is None:
        return ""
    if isinstance(valor, float) and math.isnan(valor):
        return ""
    texto = str(valor).strip().lower()
    texto = " ".join(texto.split())
    return texto


def montar_assinatura_linha(mes, unidade, quem, assunto, resolucao=""):
    """Cria uma chave de comparação usando os campos importantes da devolutiva."""
    campos = [mes, unidade, quem, assunto, resolucao]
    return " | ".join(normalizar_texto_comparacao(campo) for campo in campos)


def obter_valor_linha(row, indice):
    return str(row.iloc[indice]).strip() if len(row) > indice and normalizar_texto_comparacao(row.iloc[indice]) else ""


def ler_planilha_excel(caminho, aba=None):
    """Lê .xls/.xlsx; quando aba fica vazia, junta todas as abas do arquivo."""
    aba_limpa = aba.strip() if isinstance(aba, str) else aba
    sheet_name = aba_limpa if aba_limpa else None
    dados = pd.read_excel(caminho, sheet_name=sheet_name, header=None).fillna("")
    if isinstance(dados, dict):
        partes = []
        for nome_aba, df in dados.items():
            if df.empty:
                continue
            df = df.copy()
            df["__aba_origem__"] = nome_aba
            partes.append(df)
        if not partes:
            return pd.DataFrame()
        return pd.concat(partes, ignore_index=True)
    return dados

class AssistenteTriagem:
    def __init__(self, root):
        self.root = root
        self.root.title("Estúdio PRO - Comparador 100% Estável")
        
        try:
            self.root.state('zoomed')
        except tk.TclError:
            self.root.geometry("1200x800") 
            
        style = ttk.Style()
        style.theme_use('clam')
        
        self.nome_arquivo_atual = ""
        self.id_historico_atual = ""  
        self.ultimo_indice_busca = "1.0"
        
        self.linha_em_edicao = None
        self.janela_hist = None
        self.tree_hist = None
        
        self.poly_coords = [] 
        self.drag_mode = "new"
        self.last_x = 0
        self.last_y = 0
        
        self.arquivo_historico = "historico_status_arquivos.json"
        self.arquivo_config = "config_estudio.json" 
        self.arquivo_notas = "notas_arquivos_indecifraveis.txt" 
        
        self.arquivos_processados = {}
        self.caminho_master = "Base_Reclamacoes.xlsx"
        self.aba_padrao = "2024"
        self.abas_abertas = {}
        
        self.carregar_config()
        self.carregar_historico()

        self.root.attributes('-topmost', True)

        # ==============================================================================
        # TOPO (CONTROLES E CONFIGURAÇÕES)
        # ==============================================================================
        frame_topo = tk.Frame(self.root, bg="#2c3e50", pady=10, padx=15)
        frame_topo.pack(fill=tk.X)
        
        tk.Label(frame_topo, text="Planilha Master:", bg="#2c3e50", fg="white", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=5)
        self.entry_arquivo_saida = tk.Entry(frame_topo, width=40, font=("Segoe UI", 10))
        self.entry_arquivo_saida.insert(0, self.caminho_master)
        self.entry_arquivo_saida.pack(side=tk.LEFT, padx=(5, 0))
        
        btn_procurar_master = tk.Button(frame_topo, text="📂", bg="#f39c12", fg="white", font=("Segoe UI", 9, "bold"), command=self.escolher_master)
        btn_procurar_master.pack(side=tk.LEFT, padx=(0, 15))
        
        tk.Label(frame_topo, text="Aba/Ano:", bg="#2c3e50", fg="#f1c40f", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=2)
        self.entry_aba = tk.Entry(frame_topo, width=8, font=("Segoe UI", 11, "bold"))
        self.entry_aba.insert(0, self.aba_padrao)
        self.entry_aba.pack(side=tk.LEFT, padx=5)

        btn_comparar = tk.Button(frame_topo, text="⚖️ COMPARAR PLANILHAS", bg="#3498db", fg="white", font=("Segoe UI", 9, "bold"), command=self.abrir_comparador)
        btn_comparar.pack(side=tk.RIGHT, padx=5)

        btn_historico = tk.Button(frame_topo, text="📜 VER HISTÓRICO", bg="#8e44ad", fg="white", font=("Segoe UI", 9, "bold"), command=self.mostrar_historico)
        btn_historico.pack(side=tk.RIGHT, padx=5)

        btn_fechar = tk.Button(frame_topo, text="❌ FECHAR ABA ATUAL", bg="#e74c3c", fg="white", font=("Segoe UI", 9, "bold"), command=self.fechar_aba_atual)
        btn_fechar.pack(side=tk.RIGHT, padx=5)

        btn_carregar = tk.Button(frame_topo, text="📁 ABRIR ARQUIVO", bg="#27ae60", fg="white", font=("Segoe UI", 9, "bold"), command=self.carregar_arquivo)
        btn_carregar.pack(side=tk.RIGHT, padx=5)

        # ==============================================================================
        # CORPO PRINCIPAL
        # ==============================================================================
        paned_window = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        paned_window.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)
        
        frame_esq = tk.LabelFrame(paned_window, text=" 👁️ ARQUIVOS ABERTOS ", font=("Segoe UI", 11, "bold"), padx=10, pady=5)
        paned_window.add(frame_esq, weight=6) 
        
        self.lbl_foco = tk.Label(frame_esq, text="Abra um arquivo. Use o mouse para desenhar e ajustar as pontas.", fg="#e67e22", font=("Segoe UI", 11, "bold"))
        self.lbl_foco.pack(anchor="w", pady=(0, 5))

        frame_busca = tk.Frame(frame_esq, bg="#ecf0f1", pady=6, padx=10)
        frame_busca.pack(fill=tk.X, side=tk.BOTTOM, pady=(5, 0))
        
        tk.Label(frame_busca, text="🔍 Buscar palavra:", font=("Segoe UI", 10, "bold"), bg="#ecf0f1", fg="#2c3e50").pack(side=tk.LEFT, padx=(0, 5))
        self.entry_busca = tk.Entry(frame_busca, font=("Segoe UI", 10), width=22)
        self.entry_busca.pack(side=tk.LEFT, padx=5)
        self.entry_busca.bind("<Return>", lambda e: self.executar_busca()) 
        
        tk.Button(frame_busca, text="Buscar", bg="#3498db", fg="white", font=("Segoe UI", 9, "bold"), command=self.executar_busca).pack(side=tk.LEFT, padx=2)
        tk.Button(frame_busca, text="Próxima", bg="#2ecc71", fg="white", font=("Segoe UI", 9, "bold"), command=self.proxima_busca).pack(side=tk.LEFT, padx=2)
        tk.Button(frame_busca, text="Limpar", bg="#95a5a6", fg="white", font=("Segoe UI", 9), command=self.limpar_busca).pack(side=tk.LEFT, padx=2)
        
        self.lbl_resultado_busca = tk.Label(frame_busca, text="", font=("Segoe UI", 9, "bold"), bg="#ecf0f1")
        self.lbl_resultado_busca.pack(side=tk.LEFT, padx=10)

        self.notebook = ttk.Notebook(frame_esq)
        self.notebook.pack(fill=tk.BOTH, expand=True)
        self.notebook.bind("<<NotebookTabChanged>>", self.ao_trocar_aba)
        
        frame_dir = tk.LabelFrame(paned_window, text=" 📋 MONTAR LINHA DE DESTINO ", font=("Segoe UI", 11, "bold"), padx=15, pady=5)
        paned_window.add(frame_dir, weight=4)
        
        self.campo_mes = self.criar_campo_simples(frame_dir, "MÊS:")
        self.campo_unidade = self.criar_campo_simples(frame_dir, "UNIDADE:")
        self.campo_quem = self.criar_campo_simples(frame_dir, "QUEM (Setor):")
        self.campo_quem.insert(0, "DAS") 
        
        frame_assunto = tk.Frame(frame_dir)
        frame_assunto.pack(fill=tk.BOTH, expand=True, pady=(5, 0))
        frame_ctrl_assunto = tk.Frame(frame_assunto)
        frame_ctrl_assunto.pack(fill=tk.X)
        tk.Label(frame_ctrl_assunto, text="ASSUNTO DA RECLAMAÇÃO:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        tk.Button(frame_ctrl_assunto, text="📋 COLAR", bg="#f39c12", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.colar_para_texto_longo(self.campo_assunto)).pack(side=tk.RIGHT, padx=2)
        tk.Button(frame_ctrl_assunto, text="➕ JUNTAR", bg="#9b59b6", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.puxar_para_texto_longo(self.campo_assunto, limpar_antes=False)).pack(side=tk.RIGHT, padx=2)
        tk.Button(frame_ctrl_assunto, text="⬅️ PUXAR NOVO", bg="#3498db", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.puxar_para_texto_longo(self.campo_assunto, limpar_antes=True)).pack(side=tk.RIGHT, padx=2)
        self.campo_assunto = scrolledtext.ScrolledText(frame_assunto, height=4, wrap=tk.WORD, font=("Segoe UI", 11))
        self.campo_assunto.pack(fill=tk.BOTH, expand=True, pady=(2, 5))

        frame_resolucao = tk.Frame(frame_dir)
        frame_resolucao.pack(fill=tk.BOTH, expand=True, pady=(5, 0))
        frame_ctrl_resolucao = tk.Frame(frame_resolucao)
        frame_ctrl_resolucao.pack(fill=tk.X)
        tk.Label(frame_ctrl_resolucao, text="RESOLUÇÃO / DEVOLUTIVA:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        tk.Button(frame_ctrl_resolucao, text="📋 COLAR", bg="#f39c12", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.colar_para_texto_longo(self.campo_resolucao)).pack(side=tk.RIGHT, padx=2)
        tk.Button(frame_ctrl_resolucao, text="➕ JUNTAR", bg="#9b59b6", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.puxar_para_texto_longo(self.campo_resolucao, limpar_antes=False)).pack(side=tk.RIGHT, padx=2)
        tk.Button(frame_ctrl_resolucao, text="⬅️ PUXAR NOVO", bg="#3498db", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.puxar_para_texto_longo(self.campo_resolucao, limpar_antes=True)).pack(side=tk.RIGHT, padx=2)
        self.campo_resolucao = scrolledtext.ScrolledText(frame_resolucao, height=4, wrap=tk.WORD, font=("Segoe UI", 11))
        self.campo_resolucao.pack(fill=tk.BOTH, expand=True, pady=(2, 5))
        
        frame_botoes = tk.Frame(frame_dir)
        frame_botoes.pack(fill=tk.X, side=tk.BOTTOM, pady=(5, 0))
        tk.Frame(frame_botoes, height=2, bg="#bdc3c7").pack(fill=tk.X, pady=(0, 5)) 
        
        self.lbl_modo_edicao = tk.Label(frame_botoes, text="", fg="#c0392b", font=("Segoe UI", 9, "bold"))
        self.lbl_modo_edicao.pack(fill=tk.X)

        btn_buscar_master = tk.Button(frame_botoes, text="🔎 Buscar / Editar Reclamação Existente", bg="#34495e", fg="white", font=("Segoe UI", 10, "bold"), command=self.abrir_busca_master)
        btn_buscar_master.pack(fill=tk.X, pady=(0, 2))

        self.btn_salvar = tk.Button(frame_botoes, text="💾 INSERIR NA PLANILHA MASTER", bg="#e67e22", fg="white", font=("Segoe UI", 12, "bold"), height=2, command=self.salvar_registro)
        self.btn_salvar.pack(fill=tk.X, pady=2)
        
        btn_limpar = tk.Button(frame_botoes, text="🧹 Limpar Campos (Cancelar Edição / Exceto Mês)", font=("Segoe UI", 9), command=self.limpar_campos)
        btn_limpar.pack(fill=tk.X, pady=(2, 5))

        btn_indecifravel = tk.Button(frame_botoes, text="❓ Marcar Arquivo como Indecifrável/Dúvida", bg="#7f8c8d", fg="white", font=("Segoe UI", 9, "bold"), command=self.marcar_indecifravel)
        btn_indecifravel.pack(fill=tk.X, pady=(0, 5))

        self.var_topmost = tk.BooleanVar(value=True)
        chk_topmost = tk.Checkbutton(frame_dir, text="Manter Estúdio sempre por cima das janelas", variable=self.var_topmost, command=self.toggle_topmost, font=("Segoe UI", 9))
        chk_topmost.pack(pady=(2,0))

    # ==============================================================================
    # ⚖️ COMPARADOR DE PLANILHAS BLINDADO (FIM DOS TRAVAMENTOS E ESCONDIDOS)
    # ==============================================================================
    def abrir_comparador(self):
        janela_comp = tk.Toplevel(self.root)
        janela_comp.title("⚖️ Comparador Inteligente de Planilhas")
        janela_comp.geometry("1050x550")
        # Removi o attributes('-topmost') para garantir que os popups de erro nunca se escondam atrás dele!
        
        frame_top = tk.Frame(janela_comp, bg="#2c3e50", pady=15, padx=15)
        frame_top.pack(side=tk.TOP, fill=tk.X)
        tk.Label(frame_top, text="Selecione as planilhas para comparar o conteúdo e evitar duplicatas:", font=("Segoe UI", 11, "bold"), bg="#2c3e50", fg="white").pack(anchor="w")

        frame_inputs = tk.Frame(frame_top, bg="#2c3e50")
        frame_inputs.pack(fill=tk.X, pady=5)
        
        tk.Label(frame_inputs, text="Planilha 1 (A Nova que você recebeu):", font=("Segoe UI", 10, "bold"), bg="#2c3e50", fg="#f1c40f").grid(row=0, column=0, sticky="e", padx=5, pady=2)
        entry_p1 = tk.Entry(frame_inputs, width=60, font=("Segoe UI", 10))
        entry_p1.grid(row=0, column=1, padx=5, pady=2)
        
        def buscar_p1():
            p = filedialog.askopenfilename(parent=janela_comp, filetypes=[("Excel", "*.xlsx *.xls")])
            if p: 
                entry_p1.delete(0, tk.END)
                entry_p1.insert(0, os.path.normpath(p))
            
        tk.Button(frame_inputs, text="📂 Buscar", command=buscar_p1).grid(row=0, column=2, padx=5, pady=2)

        tk.Label(frame_inputs, text="Planilha 2 (Sua Planilha Master):", font=("Segoe UI", 10, "bold"), bg="#2c3e50", fg="#3498db").grid(row=1, column=0, sticky="e", padx=5, pady=2)
        entry_p2 = tk.Entry(frame_inputs, width=60, font=("Segoe UI", 10))
        entry_p2.insert(0, os.path.normpath(self.entry_arquivo_saida.get()))
        entry_p2.grid(row=1, column=1, padx=5, pady=2)
        
        def buscar_p2():
            p = filedialog.askopenfilename(parent=janela_comp, filetypes=[("Excel", "*.xlsx *.xls")])
            if p: 
                entry_p2.delete(0, tk.END)
                entry_p2.insert(0, os.path.normpath(p))
            
        tk.Button(frame_inputs, text="📂 Buscar", command=buscar_p2).grid(row=1, column=2, padx=5, pady=2)

        tk.Label(frame_inputs, text="Aba da Planilha 1 (vazio = todas):", font=("Segoe UI", 10, "bold"), bg="#2c3e50", fg="white").grid(row=2, column=0, sticky="e", padx=5, pady=2)
        entry_aba_p1 = tk.Entry(frame_inputs, width=15, font=("Segoe UI", 10))
        entry_aba_p1.grid(row=2, column=1, sticky="w", padx=5, pady=2)

        tk.Label(frame_inputs, text="Aba/Ano da Master (Ex: 2025):", font=("Segoe UI", 10, "bold"), bg="#2c3e50", fg="white").grid(row=3, column=0, sticky="e", padx=5, pady=2)
        entry_aba_comp = tk.Entry(frame_inputs, width=15, font=("Segoe UI", 10))
        entry_aba_comp.insert(0, self.entry_aba.get())
        entry_aba_comp.grid(row=3, column=1, sticky="w", padx=5, pady=2)

        # Fundo fixo para não sumir
        frame_bottom_area = tk.Frame(janela_comp)
        frame_bottom_area.pack(side=tk.BOTTOM, fill=tk.X, padx=15, pady=10)
        
        frame_botoes = tk.Frame(frame_bottom_area)
        frame_botoes.pack(side=tk.TOP, fill=tk.X, pady=(0, 5))
        
        lbl_status = tk.Label(frame_bottom_area, text="Status: Aguardando iniciar...", font=("Segoe UI", 10, "italic"), fg="#7f8c8d")
        lbl_status.pack(side=tk.TOP)

        # Meio (Tabela)
        colunas = ("status", "mes", "unidade", "quem", "assunto")
        tree_comp = ttk.Treeview(janela_comp, columns=colunas, show="headings", height=8)
        tree_comp.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=15, pady=(5, 0))
        
        tree_comp.heading("status", text="Análise (Status)")
        tree_comp.heading("mes", text="Mês")
        tree_comp.heading("unidade", text="Unidade")
        tree_comp.heading("quem", text="Setor")
        tree_comp.heading("assunto", text="Assunto / Conteúdo")
        
        tree_comp.column("status", width=120, anchor="center")
        tree_comp.column("mes", width=80, anchor="center")
        tree_comp.column("unidade", width=120, anchor="center")
        tree_comp.column("quem", width=100, anchor="center")
        tree_comp.column("assunto", width=450)
        
        tree_comp.tag_configure("nova", foreground="#27ae60", font=("Segoe UI", 9, "bold"))
        tree_comp.tag_configure("repetida", foreground="#c0392b", font=("Segoe UI", 9, "strike"))
        
        linhas_novas_para_exportar = []

        def analisar_planilhas():
            p1_path = os.path.normpath(entry_p1.get().strip())
            p2_path = os.path.normpath(entry_p2.get().strip())
            aba_p2 = entry_aba_comp.get().strip()
            
            if not p1_path or p1_path == "." or not os.path.exists(p1_path):
                messagebox.showerror("Erro de Arquivo", "A Planilha Nova (Planilha 1) não foi encontrada.\nVerifique o caminho.", parent=janela_comp)
                return
            if not p2_path or p2_path == "." or not os.path.exists(p2_path):
                messagebox.showerror("Erro de Arquivo", "A Planilha Master (Planilha 2) não foi encontrada.\nVerifique o caminho.", parent=janela_comp)
                return
                
            # FEEDBACK DE AÇÃO IMEDIATO! Mostra que o botão funciona.
            messagebox.showinfo("Iniciando Análise", "O cruzamento de dados vai começar agora.\n\nPor favor, aguarde alguns segundos. O botão ficará cinza enquanto o computador trabalha.", parent=janela_comp)
            
            btn_analisar.config(text="⏳ PROCESSANDO... AGUARDE", state=tk.DISABLED)
            lbl_status.config(text="Status: Iniciando leitura da Planilha Nova...", fg="#e67e22")
            janela_comp.update_idletasks() # Força a tela a atualizar imediatamente
            
            try:
                for item in tree_comp.get_children(): tree_comp.delete(item)
                linhas_novas_para_exportar.clear()
                
                # Leitura Super Segura da Planilha 1 (Sem Excel invisível travando tudo)
                df1 = None
                try:
                    df1 = ler_planilha_excel(p1_path, entry_aba_p1.get())
                except Exception as e_ler:
                    lbl_status.config(text="Status: Erro de formato no arquivo Novo.", fg="#c0392b")
                    messagebox.showerror("Ação Necessária: Formato Incompatível", 
                        "O programa não conseguiu ler a 'Planilha 1'.\n\n"
                        "Isso geralmente acontece quando sistemas corporativos geram relatórios que parecem ser '.xls', mas na verdade são protegidos.\n\n"
                        "SOLUÇÃO RÁPIDA E GARANTIDA:\n"
                        "1. Abra esse arquivo Novo no seu próprio Excel.\n"
                        "2. Vá em 'Arquivo' -> 'Salvar Como'.\n"
                        "3. Salve o documento escolhendo o tipo 'Pasta de Trabalho do Excel (*.xlsx)'.\n"
                        "4. Volte aqui, selecione esse arquivo .xlsx novo e compare!", parent=janela_comp)
                    return
                    
                lbl_status.config(text="Status: Lendo Planilha Master (Isso pode demorar dependendo do tamanho)...", fg="#2980b9")
                janela_comp.update_idletasks()

                wb = load_workbook(p2_path, data_only=True)
                if aba_p2 not in wb.sheetnames:
                    messagebox.showerror("Erro de Aba", f"A aba '{aba_p2}' não foi encontrada na sua planilha Master.", parent=janela_comp)
                    return
                ws = wb[aba_p2]
                
                dados_master = []
                assinaturas_master = set()
                assinaturas_novas_vistas = set()
                for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if i == 1 or not any(row): continue 
                    mes = row[0] if len(row)>0 else ""
                    unid = row[1] if len(row)>1 else ""
                    quem = row[2] if len(row)>2 else ""
                    ass = row[3] if len(row)>3 else ""
                    resol = row[4] if len(row)>4 else ""
                    assinatura = montar_assinatura_linha(mes, unid, quem, ass, resol)
                    dados_master.append((
                        normalizar_texto_comparacao(mes), normalizar_texto_comparacao(unid),
                        normalizar_texto_comparacao(quem), normalizar_texto_comparacao(ass),
                        normalizar_texto_comparacao(resol), assinatura
                    ))
                    assinaturas_master.add(assinatura)
                    
                lbl_status.config(text="Status: Cruzando dados para encontrar repetidas...", fg="#8e44ad")
                janela_comp.update_idletasks()

                for idx, row in df1.iterrows():
                    val_mes = obter_valor_linha(row, 0)
                    val_unid = obter_valor_linha(row, 1)
                    val_quem = obter_valor_linha(row, 2)
                    val_assunto = obter_valor_linha(row, 3)
                    val_resol = obter_valor_linha(row, 4)
                    
                    if val_mes.upper() == "MÊS" or val_unid.upper() == "UNIDADE":
                        continue
                        
                    if not val_mes and not val_unid and not val_assunto: continue
                    
                    assinatura_atual = montar_assinatura_linha(val_mes, val_unid, val_quem, val_assunto, val_resol)
                    is_dup = assinatura_atual in assinaturas_master or assinatura_atual in assinaturas_novas_vistas
                    # Se não for exatamente igual, faz uma comparação inteligente por campos principais.
                    if not is_dup:
                        n_mes = normalizar_texto_comparacao(val_mes)
                        n_unid = normalizar_texto_comparacao(val_unid)
                        n_quem = normalizar_texto_comparacao(val_quem)
                        n_assunto = normalizar_texto_comparacao(val_assunto)
                        n_resol = normalizar_texto_comparacao(val_resol)
                        for m_mes, m_unid, m_quem, m_ass, m_resol, _assinatura in dados_master:
                            campos_principais_iguais = n_mes == m_mes and n_unid == m_unid and n_quem == m_quem
                            assunto_parecido = bool(n_assunto and m_ass and difflib.SequenceMatcher(None, n_assunto, m_ass).ratio() >= 0.85)
                            resolucao_parecida = (not n_resol and not m_resol) or bool(n_resol and m_resol and difflib.SequenceMatcher(None, n_resol, m_resol).ratio() >= 0.85)
                            if campos_principais_iguais and assunto_parecido and resolucao_parecida:
                                is_dup = True
                                break
                                
                    assunto_resumo = val_assunto.replace('\n', ' ')[:80] + "..." if len(val_assunto)>80 else val_assunto.replace('\n', ' ')
                    
                    if is_dup:
                        tree_comp.insert("", tk.END, values=("🔴 REPETIDA", val_mes, val_unid, val_quem, assunto_resumo), tags=("repetida",))
                    else:
                        tree_comp.insert("", tk.END, values=("🟢 NOVA!", val_mes, val_unid, val_quem, assunto_resumo), tags=("nova",))
                        linhas_novas_para_exportar.append({
                            "MÊS": val_mes, "UNIDADE": val_unid, "QUEM": val_quem, 
                            "ASSUNTO DA RECLAMAÇÃO": val_assunto, "RESOLUÇÃO DA RECLAMAÇÃO": val_resol, "TAG": ""
                        })
                        assinaturas_novas_vistas.add(assinatura_atual)
                        
                lbl_status.config(text=f"Status: Análise concluída! Encontramos {len(linhas_novas_para_exportar)} reclamações inéditas.", fg="#27ae60")
                messagebox.showinfo("✅ Concluído", f"Análise Finalizada!\n\nForam encontradas {len(linhas_novas_para_exportar)} reclamações NOVAS que não estão na sua Master.", parent=janela_comp)

            except Exception as e:
                lbl_status.config(text="Status: Ocorreu um erro.", fg="#c0392b")
                messagebox.showerror("Erro Crítico", f"Ocorreu um erro desconhecido ao comparar as planilhas:\n\n{e}", parent=janela_comp)
            finally:
                btn_analisar.config(text="🔍 ANALISAR AGORA", state=tk.NORMAL)

        def exportar_novas():
            if not linhas_novas_para_exportar:
                messagebox.showinfo("Aviso", "Não há nenhuma linha 'NOVA' (verde) na tabela para migrar.", parent=janela_comp)
                return
                
            resposta = messagebox.askyesno("Migração Automática", f"Tem certeza que deseja inserir as {len(linhas_novas_para_exportar)} linhas INÉDITAS diretamente na Aba '{entry_aba_comp.get()}' da sua Planilha Master?", parent=janela_comp)
            
            if resposta:
                arq = os.path.normpath(entry_p2.get().strip())
                aba = entry_aba_comp.get().strip()
                df_novas = pd.DataFrame(linhas_novas_para_exportar)
                
                try:
                    with pd.ExcelWriter(arq, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                        ultima_linha = writer.book[aba].max_row
                        df_novas.to_excel(writer, sheet_name=aba, startrow=ultima_linha, index=False, header=False)
                        
                    try:
                        wb = load_workbook(arq)
                        ws = wb[aba]
                        estilo_alinhamento = Alignment(horizontal='center', vertical='center', wrap_text=True)
                        for row in ws.iter_rows():
                            for cell in row: cell.alignment = estilo_alinhamento
                        wb.save(arq)
                    except: pass
                    
                    messagebox.showinfo("Sucesso Total", "A Mágica foi feita!\n\nTodas as reclamações novas foram enviadas e formatadas na sua Master com sucesso!", parent=janela_comp)
                    janela_comp.destroy()
                except PermissionError:
                    messagebox.showerror("Erro: Arquivo Aberto", "A sua planilha Master está aberta no Excel.\n\nVocê precisa fechar o Excel da Master para o programa poder transferir os dados com segurança.", parent=janela_comp)
                except Exception as e:
                    messagebox.showerror("Falha ao Transferir", f"Erro:\n{e}", parent=janela_comp)

        btn_analisar = tk.Button(frame_botoes, text="🔍 ANALISAR AGORA", bg="#3498db", fg="white", font=("Segoe UI", 10, "bold"), command=analisar_planilhas)
        btn_analisar.pack(side=tk.LEFT)
        tk.Button(frame_botoes, text="🚀 MIGRAR APENAS AS 'NOVAS' PARA A MASTER", bg="#27ae60", fg="white", font=("Segoe UI", 10, "bold"), command=exportar_novas).pack(side=tk.RIGHT)

    # ==============================================================================
    # 🔎 BUSCAR, EDITAR E EXCLUIR LINHAS NA MASTER
    # ==============================================================================
    def abrir_busca_master(self):
        arq = self.entry_arquivo_saida.get().strip()
        aba = self.entry_aba.get().strip()
        
        if not arq.endswith('.xlsx'): arq += '.xlsx'
        if not os.path.exists(arq):
            messagebox.showwarning("Aviso", "A Planilha Master ainda não existe. Salve alguma reclamação primeiro antes de buscar.")
            return
            
        try:
            wb = load_workbook(arq, data_only=True)
            if aba not in wb.sheetnames:
                messagebox.showwarning("Aviso", f"A aba '{aba}' ainda não existe na planilha.")
                return
            ws = wb[aba]
            
            dados_master = []
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i == 1: continue 
                if not any(row): continue 
                
                mes = str(row[0]) if len(row)>0 and row[0] is not None else ""
                unid = str(row[1]) if len(row)>1 and row[1] is not None else ""
                quem = str(row[2]) if len(row)>2 and row[2] is not None else ""
                assunto = str(row[3]) if len(row)>3 and row[3] is not None else ""
                resol = str(row[4]) if len(row)>4 and row[4] is not None else ""
                
                dados_master.append((i, mes, unid, quem, assunto, resol))
                
        except Exception as e:
            messagebox.showerror("Erro", f"Falha ao ler planilha. Ela está aberta no Excel?\n\n{e}")
            return
            
        janela_busca = tk.Toplevel(self.root)
        janela_busca.title("Buscar Reclamação na Planilha")
        janela_busca.geometry("1000x600")
        
        frame_top = tk.Frame(janela_busca, pady=10, padx=15)
        frame_top.pack(fill=tk.X)
        tk.Label(frame_top, text="Buscar (Assunto, Setor ou Unidade):", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        entry_filtro = tk.Entry(frame_top, width=40, font=("Segoe UI", 11))
        entry_filtro.pack(side=tk.LEFT, padx=10)
        entry_filtro.focus()
        
        colunas = ("linha", "mes", "unidade", "quem", "assunto")
        tree = ttk.Treeview(janela_busca, columns=colunas, show="headings", height=15)
        tree.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)
        
        tree.heading("linha", text="Linha")
        tree.heading("mes", text="Mês")
        tree.heading("unidade", text="Unidade")
        tree.heading("quem", text="Quem/Setor")
        tree.heading("assunto", text="Resumo do Assunto")
        
        tree.column("linha", width=50, anchor="center")
        tree.column("mes", width=80, anchor="center")
        tree.column("unidade", width=120, anchor="center")
        tree.column("quem", width=120, anchor="center")
        tree.column("assunto", width=500)
        
        def popular_tabela(filtro=""):
            for item in tree.get_children(): tree.delete(item)
            for d in dados_master:
                texto_busca = f"{d[1]} {d[2]} {d[3]} {d[4]} {d[5]}".lower()
                if filtro.lower() in texto_busca:
                    assunto_resumo = d[4].replace('\n', ' ')[:80] + "..." if len(d[4]) > 80 else d[4].replace('\n', ' ')
                    tree.insert("", tk.END, iid=str(d[0]), values=(d[0], d[1], d[2], d[3], assunto_resumo))
                    
        popular_tabela()
        entry_filtro.bind("<KeyRelease>", lambda e: popular_tabela(entry_filtro.get()))
        
        frame_botoes_busca = tk.Frame(janela_busca, pady=10)
        frame_botoes_busca.pack(fill=tk.X)

        def carregar_selecionado():
            selecionado = tree.selection()
            if not selecionado: return
            linha_id = int(selecionado[0])
            linha_dados = next((d for d in dados_master if d[0] == linha_id), None)
            
            if linha_dados:
                self.entrar_modo_edicao(linha_dados)
                janela_busca.destroy()
                
        tree.bind("<Double-1>", lambda e: carregar_selecionado())
        
        def excluir_linha_selecionada():
            selecionado = tree.selection()
            if not selecionado: 
                messagebox.showwarning("Aviso", "Selecione uma linha para excluir.", parent=janela_busca)
                return
            linha_id = int(selecionado[0])
            
            resposta = messagebox.askyesno("Excluir Linha", f"ATENÇÃO: Tem certeza que deseja apagar a Linha {linha_id} da planilha Master '{aba}'?\nIsso não pode ser desfeito!", parent=janela_busca)
            
            if resposta:
                try:
                    wb = load_workbook(arq)
                    ws = wb[aba]
                    ws.delete_rows(linha_id)
                    wb.save(arq)
                    messagebox.showinfo("Sucesso", f"Linha {linha_id} apagada da planilha Master com sucesso!", parent=janela_busca)
                    janela_busca.destroy()
                    self.abrir_busca_master()
                except Exception as e:
                    messagebox.showerror("Erro", f"Erro ao apagar linha. Verifique se a planilha está fechada no Excel.\n\nDetalhes: {e}", parent=janela_busca)

        tk.Button(frame_botoes_busca, text="✏️ Carregar Reclamação para Edição", bg="#27ae60", fg="white", font=("Segoe UI", 10, "bold"), command=carregar_selecionado).pack(side=tk.LEFT, padx=15)
        tk.Button(frame_botoes_busca, text="🗑️ Excluir Linha da Planilha", bg="#e74c3c", fg="white", font=("Segoe UI", 10, "bold"), command=excluir_linha_selecionada).pack(side=tk.RIGHT, padx=15)

    def entrar_modo_edicao(self, dados_linha):
        self.linha_em_edicao = dados_linha[0]
        self.campo_mes.delete(0, tk.END)
        self.campo_mes.insert(0, dados_linha[1])
        self.campo_unidade.delete(0, tk.END)
        self.campo_unidade.insert(0, dados_linha[2])
        self.campo_quem.delete(0, tk.END)
        self.campo_quem.insert(0, dados_linha[3])
        self.campo_assunto.delete(1.0, tk.END)
        self.campo_assunto.insert(tk.END, dados_linha[4])
        self.campo_resolucao.delete(1.0, tk.END)
        self.campo_resolucao.insert(tk.END, dados_linha[5])
        
        self.lbl_modo_edicao.config(text=f"⚠️ MODO DE EDIÇÃO: Atualizando a Linha {self.linha_em_edicao} na Aba '{self.entry_aba.get()}'")
        self.btn_salvar.config(text=f"💾 ATUALIZAR LINHA {self.linha_em_edicao} EXISTENTE", bg="#8e44ad")

    def sair_modo_edicao(self):
        self.linha_em_edicao = None
        self.lbl_modo_edicao.config(text="")
        self.btn_salvar.config(text="💾 INSERIR NA PLANILHA MASTER", bg="#e67e22")

    def limpar_campos(self):
        self.campo_unidade.delete(0, tk.END)
        self.campo_assunto.delete(1.0, tk.END)
        self.campo_resolucao.delete(1.0, tk.END)
        self.campo_quem.delete(0, tk.END)
        self.campo_quem.insert(0, "DAS")
        self.sair_modo_edicao() 

    # ==============================================================================
    # 📝 FUNÇÃO: MARCAR ARQUIVO INDECIFRÁVEL (NOTAS)
    # ==============================================================================
    def marcar_indecifravel(self):
        id_aba = self.notebook.select()
        if not id_aba or not self.id_historico_atual:
            messagebox.showwarning("Aviso", "Não há nenhum arquivo aberto no momento para ser marcado.")
            return
            
        caminho_abs = self.id_historico_atual
        nome_arq = self.nome_arquivo_atual
        
        timestamp = datetime.now().strftime("%d/%m/%Y %H:%M")
        with open(self.arquivo_notas, "a", encoding="utf-8") as f:
            f.write(f"[{timestamp}] Arquivo: {nome_arq} | Caminho: {caminho_abs}\n")
            
        if caminho_abs in self.arquivos_processados:
            self.arquivos_processados[caminho_abs]["status"] = "Indecifrável"
            self.salvar_historico()
            self.atualizar_treeview_historico() 
            
        self.lbl_foco.config(text=f"⚠️ {nome_arq} salvo nas notas de dúvida!", fg="#7f8c8d")
        messagebox.showinfo("Arquivo Separado", f"O arquivo '{nome_arq}' foi salvo no bloco de notas.\n\nO status dele no histórico também foi atualizado.")

    # ==============================================================================
    # 🎨 DESENHO E LÓGICA DO POLÍGONO INTELIGENTE
    # ==============================================================================
    def desenhar_poligono(self, canvas):
        canvas.delete("poly")
        canvas.delete("handle")
        if len(self.poly_coords) == 8:
            canvas.create_polygon(self.poly_coords, outline="red", fill="", dash=(4, 4), width=2, tags="poly")
            r = 5
            for i in range(4):
                px, py = self.poly_coords[i*2], self.poly_coords[i*2+1]
                canvas.create_rectangle(px-r, py-r, px+r, py+r, fill="#3498db", outline="white", tags="handle")

    def no_mouse_hover(self, event, canvas):
        if not self.poly_coords:
            canvas.config(cursor="crosshair")
            return
        x, y = canvas.canvasx(event.x), canvas.canvasy(event.y)
        for i in range(4):
            cx, cy = self.poly_coords[i*2], self.poly_coords[i*2+1]
            if abs(x - cx) < 12 and abs(y - cy) < 12:
                canvas.config(cursor="fleur")
                return
        min_x, min_y = min(self.poly_coords[0::2]), min(self.poly_coords[1::2])
        max_x, max_y = max(self.poly_coords[0::2]), max(self.poly_coords[1::2])
        if min_x < x < max_x and min_y < y < max_y:
            canvas.config(cursor="hand2")
        else:
            canvas.config(cursor="crosshair")

    def iniciar_selecao(self, event, canvas):
        x, y = canvas.canvasx(event.x), canvas.canvasy(event.y)
        self.drag_mode = "new"
        if self.poly_coords:
            for i in range(4):
                cx, cy = self.poly_coords[i*2], self.poly_coords[i*2+1]
                if abs(x - cx) < 12 and abs(y - cy) < 12:
                    self.drag_mode = f"corner_{i}"
                    return
            min_x, min_y = min(self.poly_coords[0::2]), min(self.poly_coords[1::2])
            max_x, max_y = max(self.poly_coords[0::2]), max(self.poly_coords[1::2])
            if min_x < x < max_x and min_y < y < max_y:
                self.drag_mode = "move"
                self.last_x, self.last_y = x, y
                return
        self.drag_mode = "new"
        self.start_x, self.start_y = x, y
        self.poly_coords = [x, y, x, y, x, y, x, y]
        self.desenhar_poligono(canvas)

    def arrastar_selecao(self, event, canvas):
        x, y = canvas.canvasx(event.x), canvas.canvasy(event.y)
        if self.drag_mode == "new":
            self.poly_coords = [self.start_x, self.start_y, x, self.start_y, x, y, self.start_x, y]
        elif self.drag_mode.startswith("corner_"):
            idx = int(self.drag_mode.split("_")[1])
            self.poly_coords[idx*2] = x
            self.poly_coords[idx*2+1] = y
        elif self.drag_mode == "move":
            dx = x - self.last_x
            dy = y - self.last_y
            for i in range(4):
                self.poly_coords[i*2] += dx
                self.poly_coords[i*2+1] += dy
            self.last_x, self.last_y = x, y
        self.desenhar_poligono(canvas)

    def finalizar_selecao(self, event, canvas):
        self.lbl_foco.config(text="Polígono ajustado! Clique em PUXAR na direita para extrair e alinhar.", fg="#27ae60")

    # ==============================================================================
    # OBTENÇÃO DINÂMICA DO NOME DA PASTA E ABERTURA DE ARQUIVOS
    # ==============================================================================
    def obter_nome_pasta(self, caminho):
        try:
            caminho_norm = os.path.normpath(caminho)
            dir_path = os.path.dirname(caminho_norm)
            parent = os.path.basename(dir_path)
            grandparent = os.path.basename(os.path.dirname(dir_path))
            
            if grandparent and not grandparent.endswith(':') and grandparent not in ['\\', '/', '']:
                return f"{grandparent} / {parent}"
            return parent or "Raiz"
        except:
            return "Raiz"

    def carregar_arquivo(self):
        caminho = filedialog.askopenfilename(title="Selecione o Arquivo", filetypes=[("Arquivos Suportados", "*.pdf *.docx *.doc *.pptx *.xlsx *.xls *.csv *.png *.jpg *.jpeg")])
        if caminho:
            self.abrir_arquivo_por_caminho(caminho)

    def abrir_arquivo_por_caminho(self, caminho):
        caminho_windows = os.path.normpath(caminho)
        
        for widget_id, dados in self.abas_abertas.items():
            if dados.get("id_hist") == caminho_windows:
                try:
                    self.notebook.select(widget_id) 
                    self.root.lift() 
                    if dados.get("tipo") == "externo":
                        try: os.startfile(caminho_windows)
                        except:
                            import subprocess
                            subprocess.Popen(['start', '', f'"{caminho_windows}"'], shell=True)
                    return
                except: pass

        extensao = caminho_windows.split('.')[-1].lower()
        nome_arquivo = os.path.basename(caminho_windows)
        pasta_pai = self.obter_nome_pasta(caminho_windows)
        
        id_hist = caminho_windows 
        nova_aba = tk.Frame(self.notebook, bg="#ecf0f1")
        self.poly_coords = [] 
        
        try:
            if id_hist in self.arquivos_processados:
                status_arq = self.arquivos_processados[id_hist]["status"]
                messagebox.showinfo("Arquivo Já Aberto", f"Você já abriu este arquivo anteriormente!\n\nArquivo: {nome_arquivo}\nStatus atual: {status_arq.upper()}")

            if extensao in ['xlsx', 'xls', 'csv', 'doc']:
                lbl_instrucoes = tk.Label(nova_aba, 
                    text=f"📊 ARQUIVO ABERTO EXTERNAMENTE!\n\n"
                         "Deixe a janela do Excel ou Word visível, selecione o texto desejado\n"
                         "e depois use os botões de 'PUXAR' aqui no painel da direita.", 
                    font=("Segoe UI", 11, "bold"), bg="#ecf0f1", fg="#2c3e50")
                lbl_instrucoes.pack(expand=True)
                
                try: os.startfile(caminho_windows)
                except:
                    import subprocess
                    subprocess.Popen(['start', '', f'"{caminho_windows}"'], shell=True)
                    
                icone = "📊" if extensao.startswith('xls') or extensao == 'csv' else "📝"
                self.notebook.add(nova_aba, text=f"{icone} {nome_arquivo[:15]}...")
                self.notebook.select(nova_aba)
                self.abas_abertas[nova_aba._w] = {"tipo": "externo", "nome": nome_arquivo, "id_hist": id_hist, "pasta": pasta_pai, "widget": None}

            elif extensao in ['pdf', 'png', 'jpg', 'jpeg']:
                paned_pdf = ttk.PanedWindow(nova_aba, orient=tk.VERTICAL)
                paned_pdf.pack(fill=tk.BOTH, expand=True)
                
                frame_pdf_imagem = tk.Frame(paned_pdf)
                paned_pdf.add(frame_pdf_imagem, weight=6) 
                canvas_pdf = self.criar_canvas_scroll(frame_pdf_imagem)
                
                canvas_pdf.bind("<Enter>", lambda e, c=canvas_pdf: c.focus_set())
                canvas_pdf.bind("<MouseWheel>", lambda e, c=canvas_pdf: c.yview_scroll(int(-1*(e.delta/120)), "units"))
                
                frame_pdf_texto = tk.LabelFrame(paned_pdf, text="Texto Backup", font=("Segoe UI", 8))
                paned_pdf.add(frame_pdf_texto, weight=1) 
                caixa_texto = scrolledtext.ScrolledText(frame_pdf_texto, wrap=tk.WORD, font=("Segoe UI", 9), bg="#ffffff", fg="#2c3e50")
                caixa_texto.pack(fill=tk.BOTH, expand=True)
                
                icone = "📄"
                if extensao in ['png', 'jpg', 'jpeg']: icone = "🖼️"
                
                if HAS_PDF_VISION:
                    images = []
                    total_h = 0
                    max_w = 0
                    self.lbl_foco.config(text="Renderizando em alta qualidade...", fg="#f39c12")
                    self.root.update()
                    
                    if extensao == 'pdf':
                        doc = fitz.open(caminho_windows)
                        for i in range(len(doc)):
                            texto_pag = doc.load_page(i).get_text("text").strip()
                            if len(texto_pag) < 15:
                                caixa_texto.insert(tk.END, f"\n[Pág {i+1}] Desenhe o polígono contornando a inclinação do texto e clique em PUXAR NOVO na direita.\n")
                            elif texto_pag: 
                                caixa_texto.insert(tk.END, f"\n--- PÁGINA {i+1} ---\n" + texto_pag + "\n")
                                
                            pix = doc.load_page(i).get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                            images.append(img)
                            total_h += img.height + 10
                            if img.width > max_w: max_w = img.width
                        doc.close()
                    else:
                        img = Image.open(caminho_windows)
                        caixa_texto.insert(tk.END, f"\nIMAGEM DETECTADA! Desenhe o polígono na imagem.\n")
                        if img.width > 2000:
                            ratio = 2000 / img.width
                            img = img.resize((int(img.width * ratio), int(img.height * ratio)), Image.LANCZOS)
                        images.append(img)
                        total_h += img.height
                        max_w = img.width
                        
                    if images:
                        comp = Image.new('RGB', (max_w, total_h), "#bdc3c7")
                        y_off = 0
                        for img in images:
                            x_off = (max_w - img.width) // 2
                            comp.paste(img, (x_off, y_off))
                            y_off += img.height + 10
                            
                        img_tk = ImageTk.PhotoImage(comp)
                        canvas_pdf.create_image(0, 0, anchor=tk.NW, image=img_tk)
                        canvas_pdf.config(scrollregion=canvas_pdf.bbox(tk.ALL))
                        canvas_pdf.image = img_tk  
                        
                        canvas_pdf.bind("<Motion>", lambda e, c=canvas_pdf: self.no_mouse_hover(e, c))
                        canvas_pdf.bind("<ButtonPress-1>", lambda e, c=canvas_pdf: self.iniciar_selecao(e, c))
                        canvas_pdf.bind("<B1-Motion>", lambda e, c=canvas_pdf: self.arrastar_selecao(e, c))
                        canvas_pdf.bind("<ButtonRelease-1>", lambda e, c=canvas_pdf: self.finalizar_selecao(e, c))

                    self.notebook.add(nova_aba, text=f"{icone} {nome_arquivo[:12]}...")
                    self.notebook.select(nova_aba)
                    self.abas_abertas[nova_aba._w] = {
                        "tipo": "pdf_imagem", 
                        "nome": nome_arquivo, 
                        "id_hist": id_hist, 
                        "pasta": pasta_pai,
                        "widget": caixa_texto,
                        "canvas": canvas_pdf,
                        "pil_image": comp if images else None
                    }
                else:
                    caixa_texto.insert(tk.END, "PyMuPDF não instalado.")

            elif extensao in ['docx', 'pptx']:
                caixa_texto = scrolledtext.ScrolledText(nova_aba, wrap=tk.WORD, font=("Segoe UI", 12), bg="#ffffff", fg="#2c3e50", padx=30, pady=20)
                caixa_texto.pack(fill=tk.BOTH, expand=True)
                
                if extensao == 'docx':
                    icone = "📝"
                    if not HAS_DOCX:
                        caixa_texto.insert(tk.END, "⚠️ ATENÇÃO: O leitor interno do Word não está instalado no seu computador.\n\npip install python-docx")
                    else:
                        try:
                            doc = docx.Document(caminho_windows)
                            leu_algo = False
                            
                            caixa_texto.insert(tk.END, "--- TEXTO DO DOCUMENTO ---\n\n")
                            for p in doc.paragraphs:
                                if p.text.strip(): 
                                    caixa_texto.insert(tk.END, p.text + "\n\n")
                                    leu_algo = True
                            
                            if doc.tables:
                                caixa_texto.insert(tk.END, "\n\n--- DADOS ENCONTRADOS EM TABELAS ---\n\n")
                                for table in doc.tables:
                                    for row in table.rows:
                                        linha_dados = []
                                        for cell in row.cells:
                                            txt = cell.text.strip().replace('\n', ' ')
                                            if txt and txt not in linha_dados: 
                                                linha_dados.append(txt)
                                        if linha_dados:
                                            caixa_texto.insert(tk.END, " | ".join(linha_dados) + "\n")
                                            leu_algo = True
                                    caixa_texto.insert(tk.END, "\n")
                                    
                            if not leu_algo:
                                caixa_texto.insert(tk.END, "⚠️ O arquivo abriu, mas o aplicativo não encontrou textos.")
                                
                        except Exception as e:
                            caixa_texto.insert(tk.END, f"⚠️ Erro ao tentar ler o arquivo Word nativamente.\n\nDetalhe: {e}")
                else:
                    icone = "📽️"
                    if not HAS_PPTX:
                        caixa_texto.insert(tk.END, "⚠️ ATENÇÃO: O leitor do PowerPoint não está instalado.\n\npip install python-pptx")
                    else:
                        try:
                            prs = pptx.Presentation(caminho_windows)
                            for i, slide in enumerate(prs.slides):
                                caixa_texto.insert(tk.END, f" 🖥️ SLIDE {i+1} ---\n")
                                for shape in slide.shapes:
                                    if hasattr(shape, "text") and shape.text.strip():
                                        caixa_texto.insert(tk.END, shape.text + "\n")
                                caixa_texto.insert(tk.END, "\n\n")
                        except Exception as e:
                            caixa_texto.insert(tk.END, f"⚠️ Erro ao ler PowerPoint: {e}")

                self.notebook.add(nova_aba, text=f"{icone} {nome_arquivo[:12]}...")
                self.notebook.select(nova_aba)
                self.abas_abertas[nova_aba._w] = {"tipo": "texto", "nome": nome_arquivo, "id_hist": id_hist, "pasta": pasta_pai, "widget": caixa_texto}

            self.campo_unidade.delete(0, tk.END)
            self.campo_unidade.insert(0, os.path.splitext(nome_arquivo)[0].upper())

            if id_hist in self.arquivos_processados:
                status_arq = self.arquivos_processados[id_hist]["status"]
                if status_arq == "Concluído": self.lbl_foco.config(text=f"✅ {nome_arquivo} (Já Concluído)", fg="#27ae60")
                elif status_arq == "Parei nesse": self.lbl_foco.config(text=f"🛑 Retomando: {nome_arquivo} (Parei neste)", fg="#c0392b")
                elif status_arq == "Não encontrado": self.lbl_foco.config(text=f"⚠️ Revisando: {nome_arquivo} (Não Encontrado)", fg="#8e44ad")
                elif status_arq == "Indecifrável": self.lbl_foco.config(text=f"⚠️ Revisando: {nome_arquivo} (Indecifrável)", fg="#7f8c8d")
                else: self.lbl_foco.config(text=f"⏳ Lendo: {nome_arquivo}", fg="#f39c12")
            else:
                self.lbl_foco.config(text=f"Lendo: {nome_arquivo}", fg="#2980b9")
                self.adicionar_ao_historico(id_hist, nome_arquivo, pasta_pai)
                
            self.root.lift()

        except Exception as e:
            messagebox.showerror("Erro de Carregamento", f"Falha ao abrir o arquivo:\n{e}")

    # ==============================================================================
    # EXTRATOR OCR COM ALINHAMENTO E SUPORTE A WORD EXTERNO (WIN32COM)
    # ==============================================================================
    def _obter_texto_foco(self):
        id_aba_atual = self.notebook.select()
        if not id_aba_atual: return ""
        dados_aba = self.abas_abertas.get(id_aba_atual)
        if not dados_aba: return ""

        if dados_aba["tipo"] == "externo":
            try:
                import win32com.client
                try:
                    excel = win32com.client.GetActiveObject("Excel.Application")
                    val = excel.ActiveCell.Value
                    if val is not None: return str(val).strip()
                except: pass
                
                try:
                    word = win32com.client.GetActiveObject("Word.Application")
                    val = word.Selection.Text
                    if val and val.strip(): return str(val).strip()
                except: pass
                
                return ""
            except: return ""
        
        elif dados_aba["tipo"] == "pdf_imagem":
            if len(self.poly_coords) == 8:
                try:
                    self.lbl_foco.config(text="Corrigindo inclinação e analisando texto...", fg="#f39c12")
                    self.root.update()
                    img_completa = dados_aba["pil_image"]
                    xs = self.poly_coords[0::2]
                    ys = self.poly_coords[1::2]
                    left, top = min(xs), min(ys)
                    right, bottom = max(xs), max(ys)
                    
                    if abs(right - left) < 10 or abs(bottom - top) < 10: return ""
                    crop_img = img_completa.crop((left, top, right, bottom))
                    mask = Image.new('L', crop_img.size, 0)
                    poly_shifted = [(self.poly_coords[i] - left, self.poly_coords[i+1] - top) for i in range(0, 8, 2)]
                    ImageDraw.Draw(mask).polygon(poly_shifted, outline=255, fill=255)
                    white_bg = Image.new('RGB', crop_img.size, (255, 255, 255))
                    masked_img = Image.composite(crop_img, white_bg, mask)
                    
                    dx = self.poly_coords[2] - self.poly_coords[0]
                    dy = self.poly_coords[3] - self.poly_coords[1]
                    angle = math.degrees(math.atan2(dy, dx))
                    
                    rotated_img = masked_img.rotate(angle, fillcolor=(255,255,255), expand=True)
                    final_img = rotated_img.convert('L')
                    
                    texto_bruto = pytesseract.image_to_string(final_img, lang='por', config='--psm 6')
                    linhas = texto_bruto.split('\n')
                    linhas_limpas = [l.strip() for l in linhas if l.strip()]
                    texto_perfeito = " ".join(linhas_limpas)
                    
                    self.lbl_foco.config(text="Extração Direta Concluída!", fg="#27ae60")
                    return texto_perfeito
                except Exception as e:
                    messagebox.showerror("Erro OCR", f"Erro no Tesseract:\n{e}")
                    return ""
            
            try: return dados_aba["widget"].selection_get().strip()
            except: 
                messagebox.showwarning("Aviso", "Desenhe o polígono em volta do texto inclinado antes de puxar.")
                return ""

        elif dados_aba["tipo"] == "texto":
            try: return dados_aba["widget"].selection_get().strip()
            except tk.TclError:
                messagebox.showwarning("Seleção Vazia", "Selecione o texto na caixa da esquerda antes de puxar.")
                return ""

    def puxar_para_simples(self, widget_destino):
        texto = self._obter_texto_foco()
        if not texto: return
        widget_destino.delete(0, tk.END)
        widget_destino.insert(0, texto)

    def puxar_para_texto_longo(self, widget_destino, limpar_antes):
        texto = self._obter_texto_foco()
        if not texto: return
        
        if limpar_antes:
            widget_destino.delete(1.0, tk.END)
            widget_destino.insert(tk.END, f"• {texto}")
        else:
            atual = widget_destino.get(1.0, tk.END).strip()
            if os.linesep in atual or atual: widget_destino.insert(tk.END, f"\n• {texto}")
            else: widget_destino.insert(tk.END, f"• {texto}")

    def ao_trocar_aba(self, event):
        id_aba_atual = self.notebook.select()
        if self.poly_coords:
            for dados in self.abas_abertas.values():
                if "canvas" in dados:
                    try: 
                        dados["canvas"].delete("poly")
                        dados["canvas"].delete("handle")
                    except: pass
            self.poly_coords = []
            
        if id_aba_atual in self.abas_abertas:
            dados = self.abas_abertas[id_aba_atual]
            self.nome_arquivo_atual = dados["nome"]
            self.id_historico_atual = dados["id_hist"]
            self.lbl_foco.config(text=f"Aba Ativa: {self.nome_arquivo_atual}", fg="#2980b9")
            self.campo_unidade.delete(0, tk.END)
            self.campo_unidade.insert(0, os.path.splitext(self.nome_arquivo_atual)[0].upper())
            self.ultimo_indice_busca = "1.0"
        else:
            self.nome_arquivo_atual = ""
            self.id_historico_atual = ""
            self.lbl_foco.config(text="Abra um ou mais arquivos...", fg="#e67e22")

    def fechar_aba_atual(self):
        id_aba_atual = self.notebook.select()
        if id_aba_atual:
            self.notebook.forget(id_aba_atual)
            del self.abas_abertas[id_aba_atual]
            self.limpar_busca()

    def criar_canvas_scroll(self, pai):
        scroll_y = tk.Scrollbar(pai)
        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        scroll_x = tk.Scrollbar(pai, orient=tk.HORIZONTAL)
        scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        canvas = tk.Canvas(pai, bg="#dcdde1", yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        canvas.pack(fill=tk.BOTH, expand=True)
        scroll_y.config(command=canvas.yview)
        scroll_x.config(command=canvas.xview)
        return canvas

    def criar_campo_simples(self, pai, texto_label):
        frame = tk.Frame(pai)
        frame.pack(fill=tk.X, pady=4)
        frame_topo = tk.Frame(frame)
        frame_topo.pack(fill=tk.X)
        tk.Label(frame_topo, text=texto_label, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        btn_colar = tk.Button(frame_topo, text="📋", bg="#f39c12", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.colar_para_simples(entrada))
        btn_colar.pack(side=tk.RIGHT, padx=(2, 0))
        btn_puxar = tk.Button(frame_topo, text="⬅️ Puxar", bg="#3498db", fg="white", font=("Segoe UI", 8, "bold"), command=lambda: self.puxar_para_simples(entrada))
        btn_puxar.pack(side=tk.RIGHT)
        entrada = tk.Entry(frame, font=("Segoe UI", 11))
        entrada.pack(fill=tk.X, pady=(2, 0))
        return entrada

    def _ler_area_transferencia(self):
        try:
            import win32clipboard
            win32clipboard.OpenClipboard()
            try: texto = win32clipboard.GetClipboardData(win32clipboard.CF_UNICODETEXT)
            except: texto = win32clipboard.GetClipboardData(win32clipboard.CF_TEXT).decode('utf-8', errors='ignore')
            win32clipboard.CloseClipboard()
            return texto
        except Exception:
            try: return self.root.clipboard_get()
            except: return ""

    def colar_para_simples(self, widget_destino):
        texto = self._ler_area_transferencia().strip()
        if texto:
            widget_destino.delete(0, tk.END)
            widget_destino.insert(0, texto)

    def colar_para_texto_longo(self, widget_destino):
        texto = self._ler_area_transferencia().strip()
        if texto:
            texto_limpo = texto.replace('\n', ' ') 
            atual = widget_destino.get(1.0, tk.END).strip()
            if os.linesep in atual or atual: widget_destino.insert(tk.END, f"\n• {texto_limpo}")
            else: widget_destino.insert(tk.END, f"• {texto_limpo}")

    def limpar_campos(self):
        self.campo_unidade.delete(0, tk.END)
        self.campo_assunto.delete(1.0, tk.END)
        self.campo_resolucao.delete(1.0, tk.END)
        self.campo_quem.delete(0, tk.END)
        self.campo_quem.insert(0, "DAS")
        self.sair_modo_edicao() 

    # ==============================================================================
    # 📡 RADAR ANTI-DUPLICATAS E SALVAMENTO SEGURO
    # ==============================================================================
    def salvar_registro(self):
        mes = self.campo_mes.get().strip()
        unidade = self.campo_unidade.get().strip()
        quem = self.campo_quem.get().strip()
        assunto = self.campo_assunto.get(1.0, tk.END).strip()
        resolucao = self.campo_resolucao.get(1.0, tk.END).strip()
        
        if not assunto and not resolucao:
            messagebox.showwarning("Aviso", "Preencha o Assunto ou a Resolução antes de guardar!")
            return

        if self.linha_em_edicao is not None:
            self.executar_salvamento()
            return

        arq = self.entry_arquivo_saida.get().strip()
        aba = self.entry_aba.get().strip()
        if not arq.endswith('.xlsx'): arq += '.xlsx'
        if not aba: aba = "Geral"

        possiveis_duplicatas = []
        try:
            if os.path.exists(arq):
                wb = load_workbook(arq, data_only=True)
                if aba in wb.sheetnames:
                    ws = wb[aba]
                    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                        if i == 1 or not any(row): continue
                        
                        mes_plan = str(row[0]) if len(row)>0 and row[0] is not None else ""
                        unid_plan = str(row[1]) if len(row)>1 and row[1] is not None else ""
                        quem_plan = str(row[2]) if len(row)>2 and row[2] is not None else ""
                        assunto_plan = str(row[3]) if len(row)>3 and row[3] is not None else ""
                        resol_plan = str(row[4]) if len(row)>4 and row[4] is not None else ""
                        
                        sim_assunto = 0
                        if assunto and assunto_plan:
                            sim_assunto = difflib.SequenceMatcher(None, assunto.lower(), assunto_plan.lower()).ratio()
                        
                        is_dup = False
                        if unidade and unid_plan and unidade.lower() == unid_plan.lower() and mes.lower() == mes_plan.lower() and quem.lower() == quem_plan.lower():
                            is_dup = True
                        if sim_assunto > 0.70:
                            is_dup = True
                            
                        if is_dup:
                            possiveis_duplicatas.append((i, mes_plan, unid_plan, quem_plan, assunto_plan, resol_plan, sim_assunto))
        except Exception: pass

        if possiveis_duplicatas:
            possiveis_duplicatas.sort(key=lambda x: x[6], reverse=True)
            melhor_match = possiveis_duplicatas[0]
            self.mostrar_alerta_duplicata(melhor_match)
        else:
            self.executar_salvamento()

    def mostrar_alerta_duplicata(self, dados_match):
        linha, mes_p, unid_p, quem_p, assunto_p, resol_p, sim = dados_match
        
        janela_aviso = tk.Toplevel(self.root)
        janela_aviso.title("⚠️ Atenção: Possível Reclamação Duplicada!")
        janela_aviso.geometry("680x380")
        
        tk.Label(janela_aviso, text="🚨 Encontramos uma reclamação parecida já salva na planilha!", font=("Segoe UI", 12, "bold"), fg="#c0392b").pack(pady=10)
        
        frame_info = tk.LabelFrame(janela_aviso, text=f" Dados Salvos na Planilha (Linha {linha}) ", font=("Segoe UI", 10, "bold"), padx=10, pady=10)
        frame_info.pack(fill=tk.BOTH, expand=True, padx=20, pady=5)
        
        tk.Label(frame_info, text=f"Unidade: {unid_p} | Quem: {quem_p} | Mês: {mes_p}", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        texto_assunto = assunto_p[:150] + "..." if len(assunto_p) > 150 else assunto_p
        tk.Label(frame_info, text=f"Assunto Original:\n{texto_assunto}", font=("Segoe UI", 10), wraplength=600, justify=tk.LEFT).pack(anchor="w", pady=5)
        
        tk.Label(janela_aviso, text="O que você deseja fazer com o que acabou de digitar?", font=("Segoe UI", 10, "bold")).pack(pady=(10, 5))
        
        frame_botoes = tk.Frame(janela_aviso)
        frame_botoes.pack(pady=5)
        
        def acao_mesclar():
            texto_atual_assunto = self.campo_assunto.get(1.0, tk.END).strip()
            texto_atual_resol = self.campo_resolucao.get(1.0, tk.END).strip()
            texto_clipboard = ""
            if texto_atual_assunto: texto_clipboard += f"COMPLEMENTO DE ASSUNTO:\n{texto_atual_assunto}\n\n"
            if texto_atual_resol: texto_clipboard += f"NOVA RESOLUÇÃO:\n{texto_atual_resol}"
            
            if texto_clipboard:
                self.root.clipboard_clear()
                self.root.clipboard_append(texto_clipboard)
                messagebox.showinfo("Texto Copiado!", "O texto que você digitou nas caixas foi salvo na sua Área de Transferência!\n\nA linha antiga foi carregada para a tela. Agora é só você clicar no botão roxo '➕ JUNTAR' para colar a sua novidade sem apagar o que já estava lá.")
            
            self.entrar_modo_edicao(dados_match[:6]) 
            janela_aviso.destroy()
            
        def acao_salvar_nova():
            janela_aviso.destroy()
            self.executar_salvamento()
            
        def acao_cancelar():
            janela_aviso.destroy()
            
        tk.Button(frame_botoes, text="✏️ Mesclar / Carregar Linha Antiga", bg="#27ae60", fg="white", font=("Segoe UI", 9, "bold"), command=acao_mesclar).pack(side=tk.LEFT, padx=10)
        tk.Button(frame_botoes, text="💾 Ignorar e Salvar como Nova", bg="#e67e22", fg="white", font=("Segoe UI", 9, "bold"), command=acao_salvar_nova).pack(side=tk.LEFT, padx=10)
        tk.Button(frame_botoes, text="❌ Cancelar", bg="#7f8c8d", fg="white", font=("Segoe UI", 9, "bold"), command=acao_cancelar).pack(side=tk.LEFT, padx=10)

    def executar_salvamento(self):
        self.salvar_config() 
        arq = self.entry_arquivo_saida.get().strip()
        if not arq.endswith('.xlsx'): arq += '.xlsx'
        aba = self.entry_aba.get().strip()
        if not aba: aba = "Geral"

        mes = self.campo_mes.get().strip()
        unidade = self.campo_unidade.get().strip()
        quem = self.campo_quem.get().strip()
        assunto = self.campo_assunto.get(1.0, tk.END).strip()
        resolucao = self.campo_resolucao.get(1.0, tk.END).strip()

        if self.linha_em_edicao is not None:
            try:
                wb = load_workbook(arq)
                ws = wb[aba]
                
                ws.cell(row=self.linha_em_edicao, column=1, value=mes)
                ws.cell(row=self.linha_em_edicao, column=2, value=unidade)
                ws.cell(row=self.linha_em_edicao, column=3, value=quem)
                ws.cell(row=self.linha_em_edicao, column=4, value=assunto)
                ws.cell(row=self.linha_em_edicao, column=5, value=resolucao)
                
                estilo_alinhamento = Alignment(horizontal='center', vertical='center', wrap_text=True)
                for col in range(1, 7):
                    ws.cell(row=self.linha_em_edicao, column=col).alignment = estilo_alinhamento
                
                wb.save(arq)
                
                if self.id_historico_atual and self.id_historico_atual in self.arquivos_processados:
                    self.arquivos_processados[self.id_historico_atual]["status"] = "Concluído"
                    self.salvar_historico()
                    self.atualizar_treeview_historico() 
                
                self.root.attributes('-topmost', False)
                messagebox.showinfo("Sucesso", f"✅ Linha {self.linha_em_edicao} atualizada com sucesso na Aba '{aba}'!")
                self.root.attributes('-topmost', True)
                
                self.sair_modo_edicao() 
                
            except Exception as e:
                self.root.attributes('-topmost', False)
                messagebox.showerror("Erro", f"Erro ao atualizar linha.\nVerifique se a planilha '{arq}' está fechada no Excel.\n\nDetalhes: {e}")
                self.root.attributes('-topmost', True)
            return

        novo_reg = pd.DataFrame([{"MÊS": mes, "UNIDADE": unidade, "QUEM": quem, "ASSUNTO DA RECLAMAÇÃO": assunto, "RESOLUÇÃO DA RECLAMAÇÃO": resolucao, "TAG": ""}])
        try:
            if os.path.exists(arq):
                with pd.ExcelWriter(arq, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                    if aba in writer.book.sheetnames:
                        ultima_linha = writer.book[aba].max_row
                        novo_reg.to_excel(writer, sheet_name=aba, startrow=ultima_linha, index=False, header=False)
                    else:
                        novo_reg.to_excel(writer, sheet_name=aba, index=False)
            else:
                with pd.ExcelWriter(arq, engine='openpyxl') as writer:
                    novo_reg.to_excel(writer, sheet_name=aba, index=False)
            
            try:
                with pd.ExcelWriter(arq, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                    ws = writer.book[aba]
                    estilo_alinhamento = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    for row in ws.iter_rows():
                        for cell in row: cell.alignment = estilo_alinhamento
                    ws.column_dimensions['A'].width = 15  
                    ws.column_dimensions['B'].width = 25  
                    ws.column_dimensions['C'].width = 25  
                    ws.column_dimensions['D'].width = 45  
                    ws.column_dimensions['E'].width = 45  
                    ws.column_dimensions['F'].width = 15  
            except: pass
            
            if self.id_historico_atual:
                if self.id_historico_atual in self.arquivos_processados:
                    self.arquivos_processados[self.id_historico_atual]["status"] = "Concluído"
                    self.salvar_historico()
                    self.atualizar_treeview_historico() 
                    self.lbl_foco.config(text=f"✅ Salvo e Concluído!", fg="#27ae60")
            
            self.root.attributes('-topmost', False)
            messagebox.showinfo("Sucesso", f"✅ Nova reclamação salva na Aba '{aba}'!")
            self.root.attributes('-topmost', True)
            
            self.campo_assunto.delete(1.0, tk.END)
            self.campo_resolucao.delete(1.0, tk.END)
            
        except PermissionError:
            self.root.attributes('-topmost', False)
            messagebox.showerror("Erro", f"A planilha Master '{arq}' está aberta no Excel.\nFeche a Master para poder salvar.")
            self.root.attributes('-topmost', True)

    # ==============================================================================
    # HISTÓRICO EM ÁRVORE SINCRONIZADO COM ABERTURA DIRETA
    # ==============================================================================
    def carregar_config(self):
        if os.path.exists(self.arquivo_config):
            try:
                with open(self.arquivo_config, 'r', encoding='utf-8') as f:
                    dados = json.load(f)
                    self.caminho_master = dados.get("master", "Base_Reclamacoes.xlsx")
                    self.aba_padrao = dados.get("aba", "2024")
            except: pass

    def salvar_config(self):
        dados = {"master": self.entry_arquivo_saida.get().strip(), "aba": self.entry_aba.get().strip()}
        try:
            with open(self.arquivo_config, 'w', encoding='utf-8') as f: json.dump(dados, f)
        except: pass

    def escolher_master(self):
        caminho = filedialog.asksaveasfilename(title="Selecione a Planilha Master", defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile=os.path.basename(self.caminho_master))
        if caminho:
            self.entry_arquivo_saida.delete(0, tk.END)
            self.entry_arquivo_saida.insert(0, caminho)
            self.salvar_config() 

    def carregar_historico(self):
        self.arquivos_processados = {}
        if os.path.exists(self.arquivo_historico):
            try:
                with open(self.arquivo_historico, "r", encoding="utf-8") as f:
                    conteudo = json.load(f)
                    for k, v in conteudo.items():
                        if isinstance(v, dict): 
                            if not k.startswith("PASTA_MANUAL"):
                                v["pasta"] = self.obter_nome_pasta(k)
                            self.arquivos_processados[k] = v
                        else: 
                            self.arquivos_processados[k] = {"nome": os.path.basename(k), "pasta": self.obter_nome_pasta(k), "status": v}
            except: pass

    def salvar_historico(self):
        try:
            with open(self.arquivo_historico, "w", encoding="utf-8") as f:
                json.dump(self.arquivos_processados, f, ensure_ascii=False, indent=4)
        except: pass

    def adicionar_ao_historico(self, id_historico, nome, pasta):
        if id_historico and id_historico not in self.arquivos_processados:
            self.arquivos_processados[id_historico] = {"nome": nome, "pasta": pasta, "status": "Em andamento"}
            self.salvar_historico()
            self.atualizar_treeview_historico()

    def adicionar_mes_manual(self):
        parent = self.janela_hist if self.janela_hist and self.janela_hist.winfo_exists() else self.root
        mes = simpledialog.askstring("Fechamento de Pasta/Mês", "Digite o Mês/Ano ou a Pasta para concluir:\n(Ex: JANEIRO 2024)", parent=parent)
        if mes:
            id_mes = f"PASTA_MANUAL_{mes.strip().upper()}"
            if id_mes not in self.arquivos_processados:
                self.arquivos_processados[id_mes] = {"nome": f"📁 {mes.strip().upper()}", "pasta": "--- FECHAMENTO ---", "status": "Concluído"}
                self.salvar_historico()
                self.atualizar_treeview_historico()

    def copiar_nao_encontrados(self):
        pendentes = []
        for path, dados in self.arquivos_processados.items():
            if dados.get("status") in ["Não encontrado", "Indecifrável"]:
                pendentes.append(f"• Status: {dados['status'].upper()} | Pasta: {dados['pasta']} | Arquivo: {dados['nome']}")
                
        if not pendentes:
            messagebox.showinfo("Aviso", "Excelente! Não há arquivos pendentes ou indecifráveis.")
            return
            
        texto_copiar = "🚨 ARQUIVOS PENDENTES / INDECIFRÁVEIS PARA REVISÃO:\n\n" + "\n".join(pendentes)
        self.root.clipboard_clear()
        self.root.clipboard_append(texto_copiar)
        messagebox.showinfo("Sucesso", "A lista foi copiada para a área de transferência!\n\nCole (Ctrl+V) no WhatsApp ou E-mail para enviar.")

    def atualizar_treeview_historico(self):
        if not self.tree_hist or not self.tree_hist.winfo_exists(): return
            
        pastas_abertas = []
        for item in self.tree_hist.get_children():
            if self.tree_hist.item(item, "open"): pastas_abertas.append(item)
                
        for item in self.tree_hist.get_children(): self.tree_hist.delete(item)
            
        pastas = set(d["pasta"] for d in self.arquivos_processados.values())
        for pasta in sorted(pastas):
            pasta_id = f"PASTA_{pasta}"
            is_open = pasta_id in pastas_abertas
            self.tree_hist.insert("", tk.END, iid=pasta_id, text=f" 📁 {pasta}", values=("",), tags=("pasta",), open=is_open)

        for caminho_abs, dados in sorted(self.arquivos_processados.items(), key=lambda x: x[1]["nome"]):
            pasta_id = f"PASTA_{dados['pasta']}"
            icone = "📄 "
            nome_arq = dados["nome"]
            if nome_arq.lower().endswith(('.png', '.jpg', '.jpeg')): icone = "🖼️ "
            elif nome_arq.lower().endswith(('.xls', '.xlsx', '.csv', '.doc')): icone = "📊 "
            elif nome_arq.startswith("📌") or nome_arq.startswith("📁"): icone = "" 
            self.tree_hist.insert(pasta_id, tk.END, iid=caminho_abs, text=f"{icone}{nome_arq}", values=(dados["status"],), tags=(dados["status"],))

    def mostrar_historico(self):
        if self.janela_hist and self.janela_hist.winfo_exists():
            self.janela_hist.lift()
            self.atualizar_treeview_historico()
            return

        self.janela_hist = tk.Toplevel(self.root)
        self.janela_hist.title("Gestor de Status: Diretórios e Arquivos")
        self.janela_hist.geometry("900x600")
        self.janela_hist.focus_force()
        
        frame_top_hist = tk.Frame(self.janela_hist)
        frame_top_hist.pack(side=tk.TOP, fill=tk.X, padx=15, pady=10)
        tk.Label(frame_top_hist, text="Árvore de Arquivos Processados:", font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT)
        tk.Button(frame_top_hist, text="➕ Adicionar Mês/Pasta", bg="#8e44ad", fg="white", font=("Segoe UI", 9, "bold"), command=self.adicionar_mes_manual).pack(side=tk.RIGHT)

        frame_bot = tk.Frame(self.janela_hist, pady=10)
        frame_bot.pack(side=tk.BOTTOM, fill=tk.X, padx=15)
        
        colunas = ("status",)
        self.tree_hist = ttk.Treeview(self.janela_hist, columns=colunas, height=12)
        self.tree_hist.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=15, pady=5)
        
        self.tree_hist.heading("#0", text="📁 Pasta / 📄 Arquivo")
        self.tree_hist.heading("status", text="Status Atual")
        self.tree_hist.column("#0", width=650)
        self.tree_hist.column("status", width=150, anchor="center")
        
        self.tree_hist.tag_configure("Concluído", foreground="#27ae60")
        self.tree_hist.tag_configure("Parei nesse", foreground="#c0392b", font=("Segoe UI", 9, "bold"))
        self.tree_hist.tag_configure("Em andamento", foreground="#2c3e50")
        self.tree_hist.tag_configure("Não encontrado", foreground="#8e44ad", font=("Segoe UI", 9, "bold"))
        self.tree_hist.tag_configure("Indecifrável", foreground="#7f8c8d", font=("Segoe UI", 9, "bold"))
        self.tree_hist.tag_configure("pasta", font=("Segoe UI", 10, "bold"), background="#ecf0f1", foreground="#2980b9")

        self.atualizar_treeview_historico()

        tk.Label(frame_bot, text="Alterar status:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        
        opcoes_status = ttk.Combobox(frame_bot, values=["Em andamento", "Parei nesse", "Concluído", "Não encontrado", "Indecifrável"], state="readonly", width=16)
        opcoes_status.pack(side=tk.LEFT, padx=10)
        
        def atualizar_status():
            selecionado = self.tree_hist.selection()
            if not selecionado: return
            id_item = selecionado[0]
            if id_item.startswith("PASTA_"):
                messagebox.showwarning("Aviso", "Abra a pasta e selecione um arquivo específico para alterar o status.")
                return
            novo_status = opcoes_status.get()
            if novo_status:
                self.arquivos_processados[id_item]["status"] = novo_status
                self.salvar_historico()
                self.atualizar_treeview_historico()

        def excluir_selecionado():
            selecionado = self.tree_hist.selection()
            if not selecionado:
                messagebox.showwarning("Aviso", "Selecione um arquivo ou pasta na lista para excluir.")
                return
            id_item = selecionado[0]
            
            if id_item.startswith("PASTA_"):
                pasta_nome = id_item.replace("PASTA_", "")
                resposta = messagebox.askyesno("Excluir Pasta", f"Tem certeza que deseja excluir a pasta '{pasta_nome}' e TODOS os arquivos dela do histórico?")
                if resposta:
                    chaves_para_remover = [k for k, v in self.arquivos_processados.items() if v["pasta"] == pasta_nome]
                    for k in chaves_para_remover: del self.arquivos_processados[k]
                    self.salvar_historico()
                    self.atualizar_treeview_historico()
                return
                
            nome_arq = self.arquivos_processados[id_item]["nome"]
            resposta = messagebox.askyesno("Confirmar Exclusão", f"Tem certeza que deseja remover '{nome_arq}' do histórico?")
            if resposta:
                del self.arquivos_processados[id_item]
                self.salvar_historico()
                self.atualizar_treeview_historico()
                if self.id_historico_atual == id_item:
                    self.lbl_foco.config(text=f"Lendo: {self.nome_arquivo_atual}", fg="#2980b9")

        def abrir_selecionado():
            selecionado = self.tree_hist.selection()
            if not selecionado: return
            id_item = selecionado[0]
            if id_item.startswith("PASTA_"): return 
            
            try:
                self.abrir_arquivo_por_caminho(id_item)
                self.root.lift()
            except Exception as e:
                messagebox.showerror("Erro", f"Não foi possível abrir o arquivo.\nEle pode ter sido apagado, movido ou a rede está desconectada.\n\nDetalhe: {e}")

        self.tree_hist.bind("<Double-1>", lambda e: abrir_selecionado())

        tk.Button(frame_bot, text="Aplicar", bg="#3498db", fg="white", font=("Segoe UI", 9, "bold"), command=atualizar_status).pack(side=tk.LEFT)
        tk.Button(frame_bot, text="🗑️ Excluir", bg="#e74c3c", fg="white", font=("Segoe UI", 9, "bold"), command=excluir_selecionado).pack(side=tk.LEFT, padx=10)
        tk.Button(frame_bot, text="📂 Abrir Arquivo", bg="#2ecc71", fg="white", font=("Segoe UI", 9, "bold"), command=abrir_selecionado).pack(side=tk.LEFT, padx=10)
        btn_relatorio = tk.Button(frame_bot, text="📋 Copiar Lista de Pendências", bg="#e67e22", fg="white", font=("Segoe UI", 9, "bold"), command=self.copiar_nao_encontrados)
        btn_relatorio.pack(side=tk.RIGHT)

    def executar_busca(self):
        self.lbl_resultado_busca.config(text="")
        id_aba_atual = self.notebook.select()
        if not id_aba_atual or id_aba_atual not in self.abas_abertas: return
        dados_aba = self.abas_abertas[id_aba_atual]
        if dados_aba["tipo"] in ["excel", "externo"]: return
        widget_txt = dados_aba["widget"]
        widget_txt.tag_remove("busca", "1.0", tk.END) 
        palavra = self.entry_busca.get().strip()
        if not palavra: return
        widget_txt.tag_configure("busca", background="#f1c40f", foreground="black") 
        idx = "1.0"
        contagem = 0
        while True:
            idx = widget_txt.search(palavra, idx, nocase=True, stopindex=tk.END)
            if not idx: break
            fim_idx = f"{idx}+{len(palavra)}c"
            widget_txt.tag_add("busca", idx, fim_idx)
            idx = fim_idx
            contagem += 1
        if contagem > 0:
            self.lbl_resultado_busca.config(text=f"📊 {contagem} res", fg="#27ae60")
            primeiro = widget_txt.search(palavra, "1.0", nocase=True, stopindex=tk.END)
            if primeiro:
                widget_txt.see(primeiro)
                self.ultimo_indice_busca = f"{primeiro}+{len(palavra)}c"
        else: self.lbl_resultado_busca.config(text="❌ 0 res", fg="#c0392b")

    def proxima_busca(self):
        id_aba_atual = self.notebook.select()
        if not id_aba_atual or id_aba_atual not in self.abas_abertas: return
        dados_aba = self.abas_abertas[id_aba_atual]
        if dados_aba["tipo"] in ["excel", "externo"]: return
        widget_txt = dados_aba["widget"]
        palavra = self.entry_busca.get().strip()
        if not palavra: return
        prox = widget_txt.search(palavra, self.ultimo_indice_busca, nocase=True, stopindex=tk.END)
        if not prox: prox = widget_txt.search(palavra, "1.0", nocase=True, stopindex=tk.END)
        if prox:
            widget_txt.see(prox)
            self.ultimo_indice_busca = f"{prox}+{len(palavra)}c"

    def limpar_busca(self):
        self.entry_busca.delete(0, tk.END)
        self.lbl_resultado_busca.config(text="")
        for id_aba in self.abas_abertas:
            dados = self.abas_abertas[id_aba]
            if dados["tipo"] not in ["excel", "externo"]:
                try: dados["widget"].tag_remove("busca", "1.0", tk.END)
                except: pass
        self.ultimo_indice_busca = "1.0"

    def toggle_topmost(self):
        self.root.attributes('-topmost', self.var_topmost.get())

if __name__ == "__main__":
    root = tk.Tk()
    app = AssistenteTriagem(root)
    root.mainloop()